#!/usr/bin/env python3
"""Join the product availability CSV to complete regional Azure Retail Prices snapshots.

    python site/tools/match_service_availability_pricing.py

The output workbook contains all source availability rows, all API records in complete
regional cache files, and the match rationale for each availability row.
"""
import argparse
import csv
import gzip
import json
import pathlib
import re
import unicodedata
import zipfile
from collections import defaultdict
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = pathlib.Path(__file__).resolve().parents[2]
DEFAULT_CSV = ROOT / "service_availability_product_sku.csv"
DEFAULT_CACHE = ROOT / "temp" / "regional_price_cache"
DEFAULT_NONREGIONAL_CACHE = ROOT / "temp" / "nonregional_price_cache"
DEFAULT_OUTPUT = ROOT / "service_availability_pricing_enriched.xlsx"
DOC_URL = "https://learn.microsoft.com/en-us/rest/api/cost-management/retail-prices/azure-retail-prices"
RESTRICTION = "NotAvailableForSubscription"

EXTRA_HEADERS = [
    "Retail pricing status",
    "Pricing match level",
    "Pricing confidence",
    "Matched API record count",
    "Matched API service names (up to 5)",
    "Matched API product names (up to 5)",
    "Matched API SKU names (up to 5)",
    "Matched ARM SKU names (up to 5)",
    "Matched units of measure (up to 10)",
    "Matched currency codes",
    "Pricing snapshot finished at",
]
NONREGIONAL_HEADERS = [header.replace("Matched", "Non-regional matched", 1) for header in EXTRA_HEADERS]
NONREGIONAL_HEADERS[0] = "Non-regional pricing status"
NONREGIONAL_HEADERS[1] = "Non-regional match level"
NONREGIONAL_HEADERS[2] = "Non-regional confidence"
NONREGIONAL_HEADERS[-1] = "Non-regional snapshot finished at"
NONREGIONAL_HEADERS.insert(4, "Non-regional armRegionName values (up to 10)")
RAW_METADATA_HEADERS = [
    "Retail record key",
    "Cache region",
    "Cache file",
    "Cache started at",
    "Cache finished at",
]


def normalize(value):
    """Normalize case, Unicode compatibility, punctuation, and whitespace for exact joins."""
    return " ".join(
        re.findall(r"[a-z0-9]+", unicodedata.normalize("NFKC", str(value or "")).casefold())
    )


def append_row(sheet, values):
    cells = []
    for value in values:
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        cells.append(cell)
    sheet.append(cells)


def append_header(sheet, values):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cells.append(cell)
    sheet.append(cells)
    sheet.row_dimensions[1].height = 30


def add_index(index, key, record_id):
    if key:
        index[key].append(record_id)


def candidates(index, keys):
    found = set()
    result = []
    for key in keys:
        for record_id in index.get(key, ()):
            if record_id not in found:
                found.add(record_id)
                result.append(record_id)
    return result


def distinct_values(records, field, limit):
    values = []
    seen = set()
    for record in records:
        value = record.get(field)
        if value is None or value == "":
            continue
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        text = str(value)
        if text not in seen:
            seen.add(text)
            values.append(text)
            if len(values) == limit:
                break
    return "; ".join(values)


def inspect_snapshots(cache_files):
    regions = set()
    api_fields = set()
    item_count = 0
    for path in cache_files:
        with gzip.open(path, "rt", encoding="utf-8") as source:
            snapshot = json.load(source)
        if not snapshot.get("complete"):
            raise ValueError(f"Incomplete Retail Prices API snapshot: {path.name}")
        region = (snapshot.get("region") or "").strip().casefold()
        if not region or path.name != f"{region}.json.gz":
            raise ValueError(f"Region metadata mismatch in {path.name}")
        regions.add(region)
        items = snapshot.get("items")
        if not isinstance(items, list):
            raise ValueError(f"Missing API items in {path.name}")
        item_count += len(items)
        for item in items:
            if not isinstance(item, dict):
                raise ValueError(f"Malformed API item in {path.name}")
            api_fields.update(item)
    return regions, sorted(api_fields), item_count


def load_nonregional(cache_dir):
    """Load records whose armRegionName is blank, Global, or a billing zone."""
    snapshots = []
    for path in sorted(cache_dir.glob("*.json.gz")) if cache_dir.is_dir() else []:
        with gzip.open(path, "rt", encoding="utf-8") as source:
            snapshot = json.load(source)
        if not snapshot.get("complete"):
            raise ValueError(f"Incomplete Retail Prices API snapshot: {path.name}")
        if not isinstance(snapshot.get("items"), list):
            raise ValueError(f"Missing API items in {path.name}")
        snapshots.append((path.name, snapshot))
    return snapshots


def worksheet_row_count(archive, path):
    count = 0
    with archive.open(path) as source:
        for _, element in ElementTree.iterparse(source, events=("end",)):
            if element.tag.endswith("}row"):
                count += 1
                element.clear()
    return count


def build_indexes(items):
    by_service = defaultdict(list)
    by_service_product = defaultdict(list)
    by_service_sku = defaultdict(list)
    by_service_product_sku = defaultdict(list)
    by_product_sku = defaultdict(list)
    by_product = defaultdict(list)
    by_sku = defaultdict(list)
    for record_id, item in enumerate(items):
        service = normalize(item.get("serviceName"))
        product = normalize(item.get("productName"))
        sku_names = {
            normalize(item.get("skuName")),
            normalize(item.get("armSkuName")),
        } - {""}
        add_index(by_service, service, record_id)
        add_index(by_product, product, record_id)
        if service and product:
            add_index(by_service_product, (service, product), record_id)
        for sku in sku_names:
            add_index(by_sku, sku, record_id)
            if service:
                add_index(by_service_sku, (service, sku), record_id)
            if product:
                add_index(by_product_sku, (product, sku), record_id)
            if service and product:
                add_index(by_service_product_sku, (service, product, sku), record_id)
    return (
        by_service,
        by_service_product,
        by_service_sku,
        by_service_product_sku,
        by_product_sku,
        by_product,
        by_sku,
    )


def loose_service_key(value):
    """Normalized service name without a leading 'Azure' or trailing 'Service(s)'."""
    words = normalize(value).split()
    if words and words[0] == "azure":
        words = words[1:]
    if words and words[-1] in {"service", "services"}:
        words = words[:-1]
    return " ".join(words)


def build_loose_indexes(items):
    by_loose_service = defaultdict(list)
    by_product_prefix = defaultdict(list)
    for record_id, item in enumerate(items):
        add_index(by_loose_service, loose_service_key(item.get("serviceName")), record_id)
        product = loose_service_key(item.get("productName"))
        words = product.split()
        for length in range(1, len(words) + 1):
            by_product_prefix[" ".join(words[:length])].append(record_id)
    return by_loose_service, by_product_prefix


def loose_match(row, loose_indexes, items):
    """Service-level fallback for non-regional prices whose API names differ from the CSV."""
    if (row.get("product_name") or "").strip() or (row.get("sku_name") or "").strip():
        return None
    key = loose_service_key(row.get("service_name"))
    if not key:
        return None
    by_loose_service, by_product_prefix = loose_indexes
    if by_loose_service.get(key):
        return "Service only (Azure/Service prefix-suffix ignored)", "Medium", by_loose_service[key]
    if by_product_prefix.get(key):
        return "Service only (API product name starts with service name)", "Low", by_product_prefix[key]
    return None


def match_row(row, indexes, items, finished_at, with_region_values=False, loose_indexes=None):
    matched = []
    result = _match_row(row, indexes, items, finished_at, matched)
    if loose_indexes is not None and result[0] == "No matching Retail Price record":
        found = loose_match(row, loose_indexes, items)
        if found:
            level, confidence, ids = found
            matched = [items[record_id] for record_id in ids]
            result = [
                "Yes",
                level,
                confidence,
                len(ids),
                distinct_values(matched, "serviceName", 5),
                distinct_values(matched, "productName", 5),
                distinct_values(matched, "skuName", 5),
                distinct_values(matched, "armSkuName", 5),
                distinct_values(matched, "unitOfMeasure", 10),
                distinct_values(matched, "currencyCode", 5),
                finished_at,
            ]
    if with_region_values:
        regions = list(dict.fromkeys((item.get("armRegionName") or "(blank)") for item in matched))
        result.insert(4, "; ".join(regions[:10]))
    return result


def _match_row(row, indexes, items, finished_at, matched_out=None):
    if (row.get("restriction") or "").strip() == RESTRICTION:
        return ["Excluded - NotAvailableForSubscription", "", "", 0, "", "", "", "", "", "", ""]

    service_name = normalize(row.get("service_name"))
    service_id = normalize(row.get("service_id"))
    service_keys = list(dict.fromkeys(key for key in (service_name, service_id) if key))
    if not service_keys:
        return ["No service identifier", "", "", 0, "", "", "", "", "", "", finished_at]

    product = normalize(row.get("product_name"))
    sku = normalize(row.get("sku_name"))
    (
        by_service,
        by_service_product,
        by_service_sku,
        by_service_product_sku,
        by_product_sku,
        by_product,
        by_sku,
    ) = indexes

    matched_ids = []
    level = ""
    confidence = ""
    if sku and product:
        matched_ids = candidates(
            by_service_product_sku,
            [(service, product, sku) for service in service_keys],
        )
        if matched_ids:
            level, confidence = "Service + product + SKU", "High"
    if not matched_ids and sku:
        matched_ids = candidates(
            by_service_sku, [(service, sku) for service in service_keys]
        )
        if matched_ids:
            level, confidence = "Service + SKU", "High"
    if not matched_ids and sku and product:
        matched_ids = candidates(by_product_sku, [(product, sku)])
        if matched_ids:
            level, confidence = "Product + SKU (service key differs)", "Medium"
    if not matched_ids and product:
        matched_ids = candidates(
            by_service_product, [(service, product) for service in service_keys]
        )
        if matched_ids:
            level, confidence = "Service + product", "Medium"
    if not matched_ids and product and not sku:
        matched_ids = candidates(by_product, [product])
        if matched_ids:
            level, confidence = "Product only (service key differs)", "Low"
    if not matched_ids and sku and not product:
        matched_ids = candidates(by_sku, [sku])
        if matched_ids:
            level, confidence = "SKU only (service key differs)", "Low"
    if not matched_ids and not product and not sku:
        matched_ids = candidates(by_service, service_keys)
        if matched_ids:
            level, confidence = "Service only", "Low"
    if not matched_ids:
        return [
            "No matching Retail Price record", "", "", 0, "", "", "", "", "", "",
            finished_at,
        ]

    matched = [items[record_id] for record_id in matched_ids]
    if matched_out is not None:
        matched_out.extend(matched)
    return [
        "Yes",
        level,
        confidence,
        len(matched_ids),
        distinct_values(matched, "serviceName", 5),
        distinct_values(matched, "productName", 5),
        distinct_values(matched, "skuName", 5),
        distinct_values(matched, "armSkuName", 5),
        distinct_values(matched, "unitOfMeasure", 10),
        distinct_values(matched, "currencyCode", 5),
        finished_at,
    ]


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--availability-csv", type=pathlib.Path, default=DEFAULT_CSV)
    parser.add_argument("--pricing-cache-dir", type=pathlib.Path, default=DEFAULT_CACHE)
    parser.add_argument("--nonregional-cache-dir", type=pathlib.Path, default=DEFAULT_NONREGIONAL_CACHE)
    parser.add_argument("--output", type=pathlib.Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    with args.availability_csv.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        source_headers = reader.fieldnames
        source_rows = list(reader)
    if not source_headers:
        raise ValueError(f"Availability CSV has no header row: {args.availability_csv}")

    cache_files = sorted(args.pricing_cache_dir.glob("*.json.gz"))
    if not cache_files:
        raise FileNotFoundError(f"No API snapshots found in {args.pricing_cache_dir}")
    cache_regions, api_fields, api_count = inspect_snapshots(cache_files)
    raw_headers = RAW_METADATA_HEADERS + api_fields

    rows_by_region = defaultdict(list)
    for row_index, row in enumerate(source_rows):
        region = (row.get("region") or "").strip().casefold()
        rows_by_region[region].append(row_index)

    nonregional = load_nonregional(args.nonregional_cache_dir)
    nr_items, nr_meta, nr_fields = [], [], set()
    for file_name, snapshot in nonregional:
        for item in snapshot["items"]:
            nr_items.append(item)
            nr_meta.append((file_name, snapshot))
            nr_fields.update(item)
    nr_fields = sorted(nr_fields)
    nr_indexes = build_indexes(nr_items)
    nr_loose = build_loose_indexes(nr_items)
    nr_finished = max((snapshot.get("finished", "") for _, snapshot in nonregional), default="")
    extra_count = len(EXTRA_HEADERS) + len(NONREGIONAL_HEADERS)

    results = [None] * len(source_rows)
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "OpenAI"
    workbook.properties.title = "Service availability with Azure Retail Pricing"
    workbook.properties.subject = "CSV availability rows enriched with Azure Retail Prices API matches"
    workbook._named_styles["Normal"].font = Font(name="Arial", size=10, color="000000")

    readme = workbook.create_sheet("Read Me")
    readme.freeze_panes = "A2"
    readme.column_dimensions["A"].width = 34
    readme.column_dimensions["B"].width = 110
    append_header(readme, ["Item", "Details"])
    readme_data = [
        ("Availability source", f"{args.availability_csv} (all source columns and rows retained)"),
        ("Retail pricing source", f"{args.pricing_cache_dir} (all complete regional API snapshots)"),
        ("Availability rows", len(source_rows)),
        ("Retail API records", api_count),
        ("Complete regional snapshots", len(cache_files)),
        ("Availability regions", len(rows_by_region)),
        ("Regions with a cached snapshot", len(cache_regions)),
        (
            "Regions without a cached snapshot",
            "; ".join(sorted(set(rows_by_region) - cache_regions)),
        ),
        (
            "Price scope",
            "Snapshots are USD, type=Consumption, filtered by armRegionName; every cached record is retained.",
        ),
        (
            "Non-regional pricing",
            f"{args.nonregional_cache_dir}: records whose armRegionName is blank, Global, or a billing zone "
            "(Zone 1-6, US Gov, ...). Kept in the Retail API Non-Regional sheet and matched into the "
            "'Non-regional' columns of Availability Pricing, independent of the row's region.",
        ),
        ("Non-regional API records", len(nr_items)),
        (
            "Important",
            "Retail price presence is not a guarantee of service/SKU availability or subscription eligibility.",
        ),
        (
            "Match rules",
            "See service_availability_pricing_matching.md for match hierarchy, exclusions, and limitations.",
        ),
        ("Microsoft documentation", DOC_URL),
    ]
    for item in readme_data:
        append_row(readme, item)

    availability = workbook.create_sheet("Availability Pricing")
    availability.freeze_panes = "A2"
    append_header(availability, source_headers + EXTRA_HEADERS + NONREGIONAL_HEADERS)
    availability.auto_filter.ref = (
        f"A1:{get_column_letter(len(source_headers) + extra_count)}{len(source_rows) + 1}"
    )
    for column in range(1, len(source_headers) + extra_count + 1):
        letter = get_column_letter(column)
        availability.column_dimensions[letter].width = 18 if column <= len(source_headers) else 32

    raw = workbook.create_sheet("Retail API Raw")
    raw.freeze_panes = "A2"
    append_header(raw, raw_headers)
    raw.auto_filter.ref = f"A1:{get_column_letter(len(raw_headers))}{api_count + 1}"
    for column in range(1, len(raw_headers) + 1):
        raw.column_dimensions[get_column_letter(column)].width = 22

    nr_raw_headers = RAW_METADATA_HEADERS + nr_fields
    nr_raw = workbook.create_sheet("Retail API Non-Regional")
    nr_raw.freeze_panes = "A2"
    append_header(nr_raw, nr_raw_headers)
    nr_raw.auto_filter.ref = f"A1:{get_column_letter(len(nr_raw_headers))}{len(nr_items) + 1}"
    for column in range(1, len(nr_raw_headers) + 1):
        nr_raw.column_dimensions[get_column_letter(column)].width = 22
    for record_id, (item, (file_name, snapshot)) in enumerate(zip(nr_items, nr_meta)):
        append_row(
            nr_raw,
            [
                f"{snapshot['label']}:{record_id + 1}",
                snapshot["region"] or "(blank)",
                file_name,
                snapshot.get("started", ""),
                snapshot.get("finished", ""),
            ]
            + [item.get(field) for field in nr_fields],
        )

    for cache_path in cache_files:
        with gzip.open(cache_path, "rt", encoding="utf-8") as source:
            snapshot = json.load(source)
        region = snapshot["region"].strip().casefold()
        items = snapshot["items"]
        finished_at = snapshot.get("finished", "")
        indexes = build_indexes(items)

        for record_id, item in enumerate(items):
            values = [
                f"{region}:{record_id + 1}",
                region,
                cache_path.name,
                snapshot.get("started", ""),
                finished_at,
            ]
            values.extend(item.get(field) for field in api_fields)
            append_row(raw, values)

        for row_index in rows_by_region.get(region, []):
            results[row_index] = match_row(source_rows[row_index], indexes, items, finished_at)

    status_counts = defaultdict(int)
    nr_status_counts = defaultdict(int)
    restricted_count = 0
    for row_index, row in enumerate(source_rows):
        region = (row.get("region") or "").strip().casefold()
        if results[row_index] is None:
            if (row.get("restriction") or "").strip() == RESTRICTION:
                results[row_index] = [
                    "Excluded - NotAvailableForSubscription", "", "", 0, "", "", "", "", "", "", ""
                ]
            elif region not in cache_regions:
                results[row_index] = [
                    "No regional API snapshot", "", "", 0, "", "", "", "", "", "", ""
                ]
            else:
                raise ValueError(f"Row {row_index + 2} was not assigned a pricing result")
        status_counts[results[row_index][0]] += 1
        restricted_count += (row.get("restriction") or "").strip() == RESTRICTION
        nr_result = match_row(
            row, nr_indexes, nr_items, nr_finished, with_region_values=True, loose_indexes=nr_loose
        )
        nr_status_counts[nr_result[0]] += 1
        append_row(
            availability,
            [row.get(header, "") for header in source_headers] + results[row_index] + nr_result,
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)

    # Verify workbook readability, serialized row counts, and exclusion behavior.
    check = load_workbook(args.output, read_only=True, data_only=True)
    if check.sheetnames != ["Read Me", "Availability Pricing", "Retail API Raw", "Retail API Non-Regional"]:
        raise AssertionError("Workbook sheets are missing or out of order")
    with zipfile.ZipFile(args.output) as archive:
        availability_rows = worksheet_row_count(archive, "xl/worksheets/sheet2.xml")
        raw_rows = worksheet_row_count(archive, "xl/worksheets/sheet3.xml")
        nr_rows = worksheet_row_count(archive, "xl/worksheets/sheet4.xml")
    if availability_rows != len(source_rows) + 1:
        raise AssertionError("Availability sheet row count differs from the CSV")
    if raw_rows != api_count + 1:
        raise AssertionError("Raw API sheet row count differs from complete cache snapshots")
    if nr_rows != len(nr_items) + 1:
        raise AssertionError("Non-regional sheet row count differs from its cache snapshots")
    if status_counts["Excluded - NotAvailableForSubscription"] != restricted_count:
        raise AssertionError("Restricted rows were not excluded consistently")
    print(
        json.dumps(
            {
                "output": str(args.output),
                "availability_rows": len(source_rows),
                "serialized_availability_rows": availability_rows - 1,
                "source_csv_columns": len(source_headers),
                "retail_api_records": api_count,
                "serialized_retail_api_records": raw_rows - 1,
                "complete_cache_files": len(cache_files),
                "availability_regions": len(rows_by_region),
                "cached_regions": len(cache_regions),
                "status_counts": dict(status_counts),
                "retail_api_nonregional_records": len(nr_items),
                "nonregional_status_counts": dict(nr_status_counts),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
